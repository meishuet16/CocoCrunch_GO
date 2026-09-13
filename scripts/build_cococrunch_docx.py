from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "cococrunch-complete-feature-api-backend-plan.md"
OUTPUT = ROOT / "docs" / "CocoCrunch Complete Feature API and Backend Plan.docx"

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)

def borders(table):
    tblPr = table._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), "4")
        e.set(qn("w:color"), "D9D9D9")
        b.append(e)
    tblPr.append(b)

def set_cell_text(cell, text, bold=False, white=False):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(text)
    r.bold = bold
    r.font.size = Pt(8.2)
    if white: r.font.color.rgb = RGBColor(255,255,255)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

def add_table(doc, rows):
    table = doc.add_table(rows=1, cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    borders(table)
    for i, value in enumerate(rows[0]):
        shade(table.rows[0].cells[i], "17365D")
        set_cell_text(table.rows[0].cells[i], value, True, True)
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    table.rows[0]._tr.get_or_add_trPr().append(header)
    for index, row in enumerate(rows[1:]):
        cells = table.add_row().cells
        for i, value in enumerate(row):
            if index % 2 == 1: shade(cells[i], "F2F6FA")
            set_cell_text(cells[i], value)
    doc.add_paragraph().paragraph_format.space_after = Pt(5)

def parse_table(lines, start):
    data=[]
    i=start
    while i < len(lines) and lines[i].startswith("|"):
        if not set(lines[i].replace("|", "").strip()) <= {"-", ":"}:
            data.append([x.strip() for x in lines[i].strip().strip("|").split("|")])
        i += 1
    return data, i

def main():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.65); section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.65); section.right_margin = Inches(0.65)
    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"; normal.font.size = Pt(9.5)
    normal.paragraph_format.space_after = Pt(5)
    for style_name, size in [("Title", 22), ("Heading 1", 15), ("Heading 2", 11)]:
        s = doc.styles[style_name]; s.font.name = "Aptos Display"; s.font.size = Pt(size); s.font.color.rgb = RGBColor(0,0,0)

    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    i=0
    while i < len(lines):
        line=lines[i].rstrip()
        if not line:
            i += 1; continue
        if line.startswith("# "):
            p=doc.add_paragraph(style="Title"); p.add_run(line[2:]); p.alignment=WD_ALIGN_PARAGRAPH.LEFT
        elif line.startswith("## "):
            doc.add_paragraph(line[3:], style="Heading 1")
        elif line.startswith("### "):
            doc.add_paragraph(line[4:], style="Heading 2")
        elif line.startswith("|"):
            rows,i = parse_table(lines,i)
            if rows: add_table(doc,rows)
            continue
        elif line.startswith("- "):
            doc.add_paragraph(line[2:], style="List Bullet")
        elif len(line)>3 and line[0].isdigit() and ". " in line[:5]:
            doc.add_paragraph(line.split(". ",1)[1], style="List Number")
        else:
            text=line.replace("**", "").replace("`", "")
            doc.add_paragraph(text)
        i += 1
    footer=section.footer.paragraphs[0]
    footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("CocoCrunch | Feature API and Backend Plan").font.size=Pt(8)
    doc.core_properties.title = "CocoCrunch Complete Feature API and Backend Plan"
    doc.save(OUTPUT)

if __name__ == "__main__": main()
